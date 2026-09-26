"""Program utama: menjalankan perbandingan kinerja & keamanan 7 skenario.

Alur kerja:

1. Uji kebenaran (correctness) seluruh skenario.
2. Ambil plaintext uji dari dokumen resmi negara (Pancasila, Proklamasi,
   Pembukaan UUD 1945, pasal-pasal UUD 1945, UU ITE, UU PDP) yang
   dinormalkan menjadi huruf A-Z, pada ukuran 50, 100, 500, 1000, 5000.
   Setiap ukuran memakai beberapa sampel berbeda untuk pengukuran waktu.
3. Ukur waktu enkripsi/dekripsi (beserta standar deviasi), IC, entropi,
   frekuensi, dan avalanche effect tiap skenario.
4. Hitung ruang kunci + estimasi waktu brute force.
5. Ekspor hasil ke CSV & Excel, buat grafik, dan cetak ringkasan naratif.

Jalankan dengan::

    python main_comparison.py

(tidak ada dependensi wajib selain pustaka standar; grafik & Excel membutuhkan
``matplotlib`` dan ``openpyxl`` bila tersedia).
"""

import csv
import math
import os

import algorithms
import evaluation as ev
import sample_data

# ---------------------------------------------------------------------------
# Konfigurasi pengujian
# ---------------------------------------------------------------------------
OUTPUT_DIR = "hasil"
SIZES = [50, 100, 500, 1000, 5000]
REFERENCE_SIZE = 1000          # ukuran acuan untuk tabel ringkasan (waktu)
SAMPLES_PER_SIZE = 3           # banyak sampel plaintext berbeda per ukuran
AVALANCHE_TRIALS = 10          # banyak posisi flip acak per pengukuran avalanche
DISPERSION_TRIALS = 10         # banyak posisi flip acak per pengukuran dispersi
DISPERSION_SEED = 12345        # seed agar hasil dispersi dapat direproduksi


def repeats_for(size):
    """Banyak pengulangan pengukuran waktu per ukuran data (30-100x)."""
    return 30 if size >= 5000 else 100


# ---------------------------------------------------------------------------
# 1. Uji kebenaran
# ---------------------------------------------------------------------------
def run_correctness():
    print("=" * 70)
    print("[1] Uji kebenaran (decrypt(encrypt(x)) == x) untuk semua skenario")
    print("=" * 70)
    algorithms.verify_scenarios()
    print()


# ---------------------------------------------------------------------------
# 2. Waktu proses
# ---------------------------------------------------------------------------
def run_time_benchmarks(plaintext_sets):
    """Ukur waktu enkripsi/dekripsi tiap skenario pada tiap ukuran (multi sampel)."""
    print("=" * 70)
    print("[2] Pengukuran waktu proses (enkripsi & dekripsi)")
    print("=" * 70)
    time_results = {}
    for scenario in algorithms.SCENARIOS:
        time_results[scenario.name] = {}
        for size in SIZES:
            samples = plaintext_sets[size]
            repeat = repeats_for(size)
            enc_means, enc_stdevs = [], []
            dec_means, dec_stdevs = [], []
            for text in samples:
                res = ev.time_encrypt_decrypt(scenario, text, repeat=repeat)
                enc_means.append(res["encrypt_ms"])
                enc_stdevs.append(res["encrypt_stdev_ms"])
                dec_means.append(res["decrypt_ms"])
                dec_stdevs.append(res["decrypt_stdev_ms"])

            counts = [repeat] * len(samples)
            enc_mean = sum(enc_means) / len(enc_means)
            dec_mean = sum(dec_means) / len(dec_means)
            enc_stdev = ev.pooled_stdev(enc_means, enc_stdevs, counts)
            dec_stdev = ev.pooled_stdev(dec_means, dec_stdevs, counts)

            time_results[scenario.name][size] = {
                "encrypt_ms": enc_mean,
                "encrypt_stdev_ms": enc_stdev,
                "decrypt_ms": dec_mean,
                "decrypt_stdev_ms": dec_stdev,
            }
            print("  {:32s} ukuran {:5d}: enc {:.4f} ± {:.4f} ms | "
                  "dec {:.4f} ± {:.4f} ms ({} sampel x {}x)".format(
                      scenario.name, size, enc_mean, enc_stdev,
                      dec_mean, dec_stdev, len(samples), repeat))
    print()
    return time_results


# ---------------------------------------------------------------------------
# 3. Metrik statistik (IC, entropi, frekuensi)
# ---------------------------------------------------------------------------
def run_statistical_metrics(plaintexts):
    """Hitung IC, entropi, dan frekuensi tiap skenario pada tiap ukuran."""
    print("=" * 70)
    print("[3] Perhitungan IC, entropi, dan frekuensi karakter")
    print("=" * 70)
    stats = {}
    for scenario in algorithms.SCENARIOS:
        stats[scenario.name] = {}
        for size in SIZES:
            body = scenario.cipher_body(plaintexts[size])
            stats[scenario.name][size] = {
                "ic": ev.index_of_coincidence(body),
                "entropy": ev.shannon_entropy(body),
                "freq": ev.frequency_analysis(body),
            }
        ref = stats[scenario.name][REFERENCE_SIZE]
        print("  {:32s} ukuran {:5d}: IC {:.4f} | entropi {:.4f} bit".format(
            scenario.name, REFERENCE_SIZE, ref["ic"], ref["entropy"]))
    print()
    return stats


# ---------------------------------------------------------------------------
# 3b. Avalanche effect
# ---------------------------------------------------------------------------
def run_avalanche(plaintexts, trials=AVALANCHE_TRIALS):
    """Ukur plaintext & key avalanche effect tiap skenario (beberapa posisi flip)."""
    print("=" * 70)
    print("[3b] Avalanche effect (plaintext & key sensitivity)")
    print("=" * 70)
    text = plaintexts[REFERENCE_SIZE]
    avalanche = {}
    for scenario in algorithms.SCENARIOS:
        p_mean, p_std = ev.avalanche_effect_stats(scenario, text, trials=trials)
        k_mean, k_std = ev.key_avalanche_effect_stats(scenario, text,
                                                      trials=trials)
        avalanche[scenario.name] = {
            "plaintext_mean": p_mean,
            "plaintext_stdev": p_std,
            "key_mean": k_mean,
            "key_stdev": k_std,
        }
        print("  {:32s}: plaintext {:.2f} ± {:.2f}% | kunci {:.2f} ± {:.2f}% "
              "({} percobaan)".format(scenario.name, p_mean, p_std,
                                      k_mean, k_std, trials))
    print()
    return avalanche


# ---------------------------------------------------------------------------
# 3c. Dispersi posisi (position dispersion)
# ---------------------------------------------------------------------------
def run_dispersion(plaintexts, trials=DISPERSION_TRIALS, seed=DISPERSION_SEED):
    """Ukur dispersi posisi perubahan ciphertext tiap skenario."""
    print("=" * 70)
    print("[3c] Dispersi posisi (position dispersion)")
    print("=" * 70)
    text = plaintexts[REFERENCE_SIZE]
    dispersion = {}
    for scenario in algorithms.SCENARIOS:
        res = ev.dispersion_effect_stats(scenario, text, trials=trials,
                                         seed=seed)
        dispersion[scenario.name] = res
        print("  {:32s}: stdev {:.2f} ± {:.2f} | mean gap {:.2f} ± {:.2f} "
              "({} percobaan)".format(
                  scenario.name, res["stdev_mean"], res["stdev_stdev"],
                  res["mean_gap_mean"], res["mean_gap_stdev"], trials))
    print()
    return dispersion


# ---------------------------------------------------------------------------
# 4. Ruang kunci
# ---------------------------------------------------------------------------
def run_key_space():
    """Hitung ruang kunci dan estimasi brute force untuk tiap skenario."""
    print("=" * 70)
    print("[4] Analisis ruang kunci & estimasi brute force")
    print("=" * 70)
    key_spaces = {}
    vlen = len(algorithms.VIGENERE_KEY)
    tlen = len(algorithms.TRANSPOSITION_KEY)
    hsize = algorithms.HILL_BLOCK_SIZE
    for scenario in algorithms.SCENARIOS:
        parts = ev.scenario_key_space(scenario, vlen, tlen, hsize)
        total = parts["total"]
        key_spaces[scenario.name] = {
            "total": total,
            "bits": math.log2(total),
            "parts": {k: v for k, v in parts.items() if k != "total"},
            "human": ev.format_key_space(total),
            "bruteforce": ev.brute_force_time(total),
        }
        print("  {:32s}: ruang kunci {} | brute force {}".format(
            scenario.name, ev.format_key_space(total),
            key_spaces[scenario.name]["bruteforce"]["human"]))
    print()
    return key_spaces


# ---------------------------------------------------------------------------
# Ekspor hasil
# ---------------------------------------------------------------------------
def _scenario_rows(time_results, stats, key_spaces, plaintext_metrics,
                   avalanche, dispersion):
    """Susun baris tabel ringkasan (baris = skenario, kolom = metrik)."""
    rows = []
    for scenario in algorithms.SCENARIOS:
        name = scenario.name
        t = time_results[name][REFERENCE_SIZE]
        s = stats[name][REFERENCE_SIZE]
        ks = key_spaces[name]
        av = avalanche[name]
        dp = dispersion[name]
        rows.append({
            "Skenario": name,
            "Waktu Enkripsi (ms)": round(t["encrypt_ms"], 4),
            "Stdev Enkripsi (ms)": round(t["encrypt_stdev_ms"], 4),
            "Waktu Dekripsi (ms)": round(t["decrypt_ms"], 4),
            "Stdev Dekripsi (ms)": round(t["decrypt_stdev_ms"], 4),
            "Throughput Enkripsi (kar/ms)":
                round(ev.throughput(REFERENCE_SIZE, t["encrypt_ms"]), 2),
            "Avalanche Plaintext (%)": round(av["plaintext_mean"], 2),
            "Avalanche Plaintext Stdev (%)": round(av["plaintext_stdev"], 2),
            "Avalanche Kunci (%)": round(av["key_mean"], 2),
            "Avalanche Kunci Stdev (%)": round(av["key_stdev"], 2),
            "Dispersi Posisi (stdev)": round(dp["stdev_mean"], 2),
            "Dispersi Posisi (mean gap)": round(dp["mean_gap_mean"], 2),
            "IC Plaintext": round(plaintext_metrics["ic"], 4),
            "IC Ciphertext": round(s["ic"], 4),
            "Entropi Plaintext (bit)": round(plaintext_metrics["entropy"], 4),
            "Entropi Ciphertext (bit)": round(s["entropy"], 4),
            "Entropi Maksimum (bit)": round(ev.MAX_ENTROPY, 4),
            "Ruang Kunci": ks["total"],
            "Ruang Kunci (bit)": round(ks["bits"], 2),
            "Estimasi Brute Force": ks["bruteforce"]["human"],
        })
    return rows


def export_csvs(rows, time_results, stats, key_spaces, avalanche, dispersion):
    """Tulis hasil ke beberapa berkas CSV."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # 1) Tabel ringkasan
    with open(os.path.join(OUTPUT_DIR, "hasil_perbandingan.csv"), "w",
              newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    # 2) Waktu per ukuran data (format panjang, dengan stdev)
    with open(os.path.join(OUTPUT_DIR, "waktu_per_ukuran.csv"), "w",
              newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Skenario", "Ukuran (karakter)",
                         "Waktu Enkripsi (ms)", "Stdev Enkripsi (ms)",
                         "Waktu Dekripsi (ms)", "Stdev Dekripsi (ms)"])
        for scenario in algorithms.SCENARIOS:
            for size in SIZES:
                t = time_results[scenario.name][size]
                writer.writerow([scenario.name, size,
                                 round(t["encrypt_ms"], 6),
                                 round(t["encrypt_stdev_ms"], 6),
                                 round(t["decrypt_ms"], 6),
                                 round(t["decrypt_stdev_ms"], 6)])

    # 3) Frekuensi karakter A-Z per skenario (pada ukuran terbesar)
    largest = SIZES[-1]
    with open(os.path.join(OUTPUT_DIR, "frekuensi_karakter.csv"), "w",
              newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Skenario"] + list(ev.ALPHABET))
        for scenario in algorithms.SCENARIOS:
            freq = stats[scenario.name][largest]["freq"]
            writer.writerow([scenario.name] +
                            [freq[ch] for ch in ev.ALPHABET])

    # 4) Rincian ruang kunci
    with open(os.path.join(OUTPUT_DIR, "ruang_kunci.csv"), "w",
              newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Skenario", "Ruang Kunci", "Ruang Kunci (bit)",
                         "Estimasi Brute Force (10^9/detik)"])
        for scenario in algorithms.SCENARIOS:
            ks = key_spaces[scenario.name]
            writer.writerow([scenario.name, ks["total"], round(ks["bits"], 2),
                             ks["bruteforce"]["human"]])

    # 5) Avalanche effect
    with open(os.path.join(OUTPUT_DIR, "avalanche.csv"), "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Skenario", "Avalanche Plaintext (%)",
                         "Stdev Plaintext (%)", "Avalanche Kunci (%)",
                         "Stdev Kunci (%)"])
        for scenario in algorithms.SCENARIOS:
            av = avalanche[scenario.name]
            writer.writerow([scenario.name,
                             round(av["plaintext_mean"], 4),
                             round(av["plaintext_stdev"], 4),
                             round(av["key_mean"], 4),
                             round(av["key_stdev"], 4)])

    # 6) Dispersi posisi
    with open(os.path.join(OUTPUT_DIR, "dispersi.csv"), "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Skenario", "Dispersi Stdev Posisi",
                         "Stdev dari Stdev", "Mean Gap Posisi",
                         "Stdev dari Mean Gap", "Percobaan"])
        for scenario in algorithms.SCENARIOS:
            dp = dispersion[scenario.name]
            writer.writerow([scenario.name,
                             round(dp["stdev_mean"], 4),
                             round(dp["stdev_stdev"], 4),
                             round(dp["mean_gap_mean"], 4),
                             round(dp["mean_gap_stdev"], 4),
                             dp["trials"]])

    print("Hasil CSV ditulis ke folder '{}':".format(OUTPUT_DIR))
    print("  - hasil_perbandingan.csv")
    print("  - waktu_per_ukuran.csv")
    print("  - frekuensi_karakter.csv")
    print("  - ruang_kunci.csv")
    print("  - avalanche.csv")
    print("  - dispersi.csv")


def export_excel(rows, time_results, stats, key_spaces, avalanche, dispersion):
    """Tulis hasil ke berkas Excel multi-sheet (jika openpyxl tersedia)."""
    try:
        from openpyxl import Workbook
    except ImportError:
        print("  (openpyxl tidak tersedia; lewati ekspor Excel)")
        return

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    wb = Workbook()

    # Sheet 1: ringkasan
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append(list(rows[0].keys()))
    for r in rows:
        ws.append(list(r.values()))

    # Sheet 2: waktu per ukuran
    ws2 = wb.create_sheet("Waktu per Ukuran")
    ws2.append(["Skenario", "Ukuran (karakter)",
                "Waktu Enkripsi (ms)", "Stdev Enkripsi (ms)",
                "Waktu Dekripsi (ms)", "Stdev Dekripsi (ms)"])
    for scenario in algorithms.SCENARIOS:
        for size in SIZES:
            t = time_results[scenario.name][size]
            ws2.append([scenario.name, size,
                        round(t["encrypt_ms"], 6),
                        round(t["encrypt_stdev_ms"], 6),
                        round(t["decrypt_ms"], 6),
                        round(t["decrypt_stdev_ms"], 6)])

    # Sheet 3: frekuensi
    largest = SIZES[-1]
    ws3 = wb.create_sheet("Frekuensi")
    ws3.append(["Skenario"] + list(ev.ALPHABET))
    for scenario in algorithms.SCENARIOS:
        freq = stats[scenario.name][largest]["freq"]
        ws3.append([scenario.name] + [freq[ch] for ch in ev.ALPHABET])

    # Sheet 4: ruang kunci
    ws4 = wb.create_sheet("Ruang Kunci")
    ws4.append(["Skenario", "Ruang Kunci", "Ruang Kunci (bit)",
                "Estimasi Brute Force (10^9/detik)"])
    for scenario in algorithms.SCENARIOS:
        ks = key_spaces[scenario.name]
        ws4.append([scenario.name, ks["total"], round(ks["bits"], 2),
                    ks["bruteforce"]["human"]])

    # Sheet 5: avalanche effect
    ws5 = wb.create_sheet("Avalanche")
    ws5.append(["Skenario", "Avalanche Plaintext (%)",
                "Stdev Plaintext (%)", "Avalanche Kunci (%)",
                "Stdev Kunci (%)"])
    for scenario in algorithms.SCENARIOS:
        av = avalanche[scenario.name]
        ws5.append([scenario.name, round(av["plaintext_mean"], 4),
                    round(av["plaintext_stdev"], 4),
                    round(av["key_mean"], 4),
                    round(av["key_stdev"], 4)])

    # Sheet 6: dispersi posisi
    ws6 = wb.create_sheet("Dispersi")
    ws6.append(["Skenario", "Dispersi Stdev Posisi", "Stdev dari Stdev",
                "Mean Gap Posisi", "Stdev dari Mean Gap", "Percobaan"])
    for scenario in algorithms.SCENARIOS:
        dp = dispersion[scenario.name]
        ws6.append([scenario.name, round(dp["stdev_mean"], 4),
                    round(dp["stdev_stdev"], 4),
                    round(dp["mean_gap_mean"], 4),
                    round(dp["mean_gap_stdev"], 4), dp["trials"]])

    wb.save(os.path.join(OUTPUT_DIR, "hasil_perbandingan.xlsx"))
    print("  - hasil_perbandingan.xlsx")


# ---------------------------------------------------------------------------
# Grafik
# ---------------------------------------------------------------------------
def make_plots(rows, time_results, stats, key_spaces, plaintext_metrics,
               avalanche, dispersion):
    """Buat grafik perbandingan dan simpan ke folder hasil (jika matplotlib)."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        print("  (matplotlib tidak tersedia; lewati pembuatan grafik)")
        return

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    names = [s.name for s in algorithms.SCENARIOS]
    plt.rcParams["figure.dpi"] = 110

    # --- (a) Bar chart waktu enkripsi & dekripsi (ukuran acuan) ---
    enc = [time_results[n][REFERENCE_SIZE]["encrypt_ms"] for n in names]
    dec = [time_results[n][REFERENCE_SIZE]["decrypt_ms"] for n in names]
    x = range(len(names))
    width = 0.4
    fig, ax = plt.subplots(figsize=(11, 5))
    ax.bar([i - width / 2 for i in x], enc, width, label="Enkripsi", color="#4C72B0")
    ax.bar([i + width / 2 for i in x], dec, width, label="Dekripsi", color="#DD8452")
    ax.set_xticks(list(x))
    ax.set_xticklabels(names, rotation=25, ha="right", fontsize=8)
    ax.set_ylabel("Waktu (ms)")
    ax.set_title("Waktu Proses per Skenario (ukuran {} karakter)".format(
        REFERENCE_SIZE))
    ax.legend()
    fig.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, "grafik_waktu_proses.png"))
    plt.close(fig)

    # --- (b) Line chart pengaruh ukuran data terhadap waktu enkripsi ---
    fig, ax = plt.subplots(figsize=(10, 6))
    for scenario in algorithms.SCENARIOS:
        ys = [time_results[scenario.name][s]["encrypt_ms"] for s in SIZES]
        ax.plot(SIZES, ys, marker="o", label=scenario.name)
    ax.set_xscale("log")
    ax.set_xlabel("Ukuran plaintext (karakter, skala log)")
    ax.set_ylabel("Waktu enkripsi (ms)")
    ax.set_title("Pengaruh Ukuran Data terhadap Waktu Enkripsi")
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, "grafik_waktu_vs_ukuran.png"))
    plt.close(fig)

    # --- (c) Bar chart entropi ---
    ent = [stats[n][REFERENCE_SIZE]["entropy"] for n in names]
    fig, ax = plt.subplots(figsize=(11, 5))
    ax.bar(x, ent, color="#55A868")
    ax.axhline(ev.MAX_ENTROPY, color="red", linestyle="--",
               label="Entropi maksimum (%.4f)" % ev.MAX_ENTROPY)
    ax.axhline(plaintext_metrics["entropy"], color="gray", linestyle=":",
               label="Entropi plaintext (%.4f)" % plaintext_metrics["entropy"])
    for i, v in enumerate(ent):
        ax.text(i, v + 0.02, "{:.3f}".format(v), ha="center", fontsize=7)
    ax.set_xticks(list(x))
    ax.set_xticklabels(names, rotation=25, ha="right", fontsize=8)
    ax.set_ylim(0, ev.MAX_ENTROPY + 0.35)
    ax.set_ylabel("Entropi (bit)")
    ax.set_title("Entropi Shannon Ciphertext per Skenario")
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, "grafik_entropi.png"))
    plt.close(fig)

    # --- (d) Bar chart Index of Coincidence ---
    ic = [stats[n][REFERENCE_SIZE]["ic"] for n in names]
    fig, ax = plt.subplots(figsize=(11, 5))
    ax.bar(x, ic, color="#C44E52")
    ax.axhline(ev.IC_RANDOM, color="green", linestyle="--",
               label="IC acak ideal (%.4f)" % ev.IC_RANDOM)
    ax.axhline(plaintext_metrics["ic"], color="gray", linestyle=":",
               label="IC plaintext (%.4f)" % plaintext_metrics["ic"])
    for i, v in enumerate(ic):
        ax.text(i, v + 0.002, "{:.4f}".format(v), ha="center", fontsize=7)
    ax.set_xticks(list(x))
    ax.set_xticklabels(names, rotation=25, ha="right", fontsize=8)
    ax.set_ylim(0, max(ic + [plaintext_metrics["ic"], 0.08]) + 0.01)
    ax.set_ylabel("Index of Coincidence")
    ax.set_title("Index of Coincidence Ciphertext per Skenario")
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, "grafik_ic.png"))
    plt.close(fig)

    # --- (e) Bar chart ruang kunci (dalam bit, skala linear) ---
    bits = [key_spaces[n]["bits"] for n in names]
    fig, ax = plt.subplots(figsize=(11, 5))
    ax.bar(x, bits, color="#8172B3")
    for i, b in enumerate(bits):
        ax.text(i, b + 0.5, "{:.1f}".format(b), ha="center", fontsize=8)
    ax.set_xticks(list(x))
    ax.set_xticklabels(names, rotation=25, ha="right", fontsize=8)
    ax.set_ylabel("Ruang kunci (bit, log2)")
    ax.set_title("Ruang Kunci per Skenario (dalam bit)")
    fig.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, "grafik_ruang_kunci.png"))
    plt.close(fig)

    # --- (f) Histogram frekuensi karakter per skenario ---
    largest = SIZES[-1]
    nrows, ncols = 4, 2
    fig, axes = plt.subplots(nrows, ncols, figsize=(13, 13))
    letters = list(ev.ALPHABET)
    for idx, scenario in enumerate(algorithms.SCENARIOS):
        ax = axes.flat[idx]
        freq = stats[scenario.name][largest]["freq"]
        counts = [freq[ch] for ch in letters]
        ax.bar(letters, counts, color="#4C72B0")
        ax.set_title(scenario.name, fontsize=9)
        ax.tick_params(axis="x", labelsize=6)
    axes.flat[-1].set_visible(False)
    fig.suptitle("Distribusi Frekuensi Karakter A-Z (ukuran {} karakter)".format(
        largest), fontsize=13)
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    fig.savefig(os.path.join(OUTPUT_DIR, "grafik_frekuensi.png"))
    plt.close(fig)

    # --- (g) Bar chart avalanche effect ---
    av_plain = [avalanche[n]["plaintext_mean"] for n in names]
    av_key = [avalanche[n]["key_mean"] for n in names]
    av_plain_err = [avalanche[n]["plaintext_stdev"] for n in names]
    av_key_err = [avalanche[n]["key_stdev"] for n in names]
    fig, ax = plt.subplots(figsize=(11, 5))
    ax.bar([i - width / 2 for i in x], av_plain, width, yerr=av_plain_err,
           capsize=3, label="Plaintext avalanche", color="#4C72B0")
    ax.bar([i + width / 2 for i in x], av_key, width, yerr=av_key_err,
           capsize=3, label="Key avalanche", color="#DD8452")
    ax.axhline(ev.AVALANCHE_IDEAL, color="red", linestyle="--",
               label="Ideal acak (%.2f%%)" % ev.AVALANCHE_IDEAL)
    ax.set_xticks(list(x))
    ax.set_xticklabels(names, rotation=25, ha="right", fontsize=8)
    ax.set_ylabel("Avalanche effect (%)")
    ax.set_title("Avalanche Effect per Skenario (ukuran {} karakter)".format(
        REFERENCE_SIZE))
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, "grafik_avalanche.png"))
    plt.close(fig)

    # --- (h) Bar chart dispersi posisi (stdev) ---
    from matplotlib.patches import Patch
    disp_stdev = [dispersion[n]["stdev_mean"] for n in names]
    disp_stdev_err = [dispersion[n]["stdev_stdev"] for n in names]
    has_trans = ["Transposisi" in s.components for s in algorithms.SCENARIOS]
    colors = ["#C44E52" if t else "#4C72B0" for t in has_trans]
    fig, ax = plt.subplots(figsize=(11, 5))
    ax.bar(x, disp_stdev, yerr=disp_stdev_err, capsize=3, color=colors)
    for i, v in enumerate(disp_stdev):
        ax.text(i, v + disp_stdev_err[i] + 0.2, "{:.1f}".format(v),
                ha="center", fontsize=8)
    ax.set_xticks(list(x))
    ax.set_xticklabels(names, rotation=25, ha="right", fontsize=8)
    ax.set_ylabel("Standar deviasi posisi (index)")
    ax.set_title("Dispersi Posisi per Skenario (stdev, {} karakter)".format(
        REFERENCE_SIZE))
    ax.legend(handles=[
        Patch(color="#4C72B0", label="Tanpa transposisi"),
        Patch(color="#C44E52", label="Dengan transposisi"),
    ], fontsize=8)
    fig.tight_layout()
    fig.savefig(os.path.join(OUTPUT_DIR, "grafik_dispersi.png"))
    plt.close(fig)

    print("Grafik ditulis ke folder '{}':".format(OUTPUT_DIR))
    print("  - grafik_waktu_proses.png")
    print("  - grafik_waktu_vs_ukuran.png")
    print("  - grafik_entropi.png")
    print("  - grafik_ic.png")
    print("  - grafik_ruang_kunci.png")
    print("  - grafik_frekuensi.png")
    print("  - grafik_avalanche.png")
    print("  - grafik_dispersi.png")


# ---------------------------------------------------------------------------
# Ringkasan naratif
# ---------------------------------------------------------------------------
def build_summary(rows, time_results, stats, key_spaces, plaintext_metrics,
                  avalanche, dispersion):
    """Tentukan skenario unggulan pada tiap metrik + susun narasi trade-off."""
    names = [r["Skenario"] for r in rows]

    fastest_enc = min(names, key=lambda n:
                      time_results[n][REFERENCE_SIZE]["encrypt_ms"])
    fastest_dec = min(names, key=lambda n:
                      time_results[n][REFERENCE_SIZE]["decrypt_ms"])
    slowest_enc = max(names, key=lambda n:
                      time_results[n][REFERENCE_SIZE]["encrypt_ms"])

    # IC terbaik = paling dekat ke nilai acak ideal 1/26.
    best_ic = min(names, key=lambda n:
                  abs(stats[n][REFERENCE_SIZE]["ic"] - ev.IC_RANDOM))

    # Entropi terbaik = paling tinggi (paling dekat entropi maksimum).
    best_entropy = max(names, key=lambda n:
                       stats[n][REFERENCE_SIZE]["entropy"])

    best_keyspace = max(names, key=lambda n: key_spaces[n]["total"])

    # Avalanche terbaik = paling dekat ke nilai ideal (25/26).
    best_avalanche = min(
        names, key=lambda n:
        abs(avalanche[n]["plaintext_mean"] - ev.AVALANCHE_IDEAL)
    )

    lines = []
    lines.append("=" * 70)
    lines.append("RINGKASAN NARATIF HASIL PERBANDINGAN")
    lines.append("=" * 70)
    lines.append("")
    lines.append("Skenario paling unggul pada tiap metrik:")
    lines.append("  - Waktu enkripsi tercepat : {}".format(fastest_enc))
    lines.append("  - Waktu dekripsi tercepat : {}".format(fastest_dec))
    lines.append("  - Waktu enkripsi terlambat: {}".format(slowest_enc))
    lines.append("  - IC paling acak (aman)   : {} (IC {:.4f})".format(
        best_ic, stats[best_ic][REFERENCE_SIZE]["ic"]))
    lines.append("  - Entropi tertinggi        : {} ({:.4f} bit)".format(
        best_entropy, stats[best_entropy][REFERENCE_SIZE]["entropy"]))
    lines.append("  - Avalanche terbaik        : {} ({:.2f}%)".format(
        best_avalanche, avalanche[best_avalanche]["plaintext_mean"]))
    lines.append("  - Ruang kunci terbesar     : {} ({})".format(
        best_keyspace, key_spaces[best_keyspace]["human"]))
    lines.append("")
    lines.append("Trade-off keamanan vs waktu proses:")
    lines.append("  * Algoritma tunggal cepat, tetapi keamanannya terbatas:")
    lines.append("    - Transposisi Kolom hanya menyusun ulang posisi sehingga")
    lines.append("      distribusi frekuensi huruf TIDAK berubah (IC ciphertext")
    lines.append("      sama dengan IC plaintext), mudah diserang analisis")
    lines.append("      frekuensi.")
    lines.append("    - Vigenere meratakan frekuensi (IC mendekati acak), tetapi")
    lines.append("      dengan kunci pendek masih rentan terhadap serangan")
    lines.append("      Kasiski/Friedman.")
    lines.append("    - Hill menyamarkan frekuensi per blok, tetapi ruang kunci")
    lines.append("      bergantung ukuran matriks dan rentan known-plaintext.")
    lines.append("  * Kombinasi (Multiple Super Encryption) meningkatkan keamanan")
    lines.append("    (IC mendekati acak, entropi mendekati maksimum, ruang kunci")
    lines.append("    berlipat ganda), namun menambah waktu proses karena setiap")
    lines.append("    lapisan memproses ulang seluruh teks.")
    lines.append("  * Skenario kombinasi penuh (Vigenere + Hill + Transposisi)")
    lines.append("    memberikan ruang kunci gabungan terbesar (perkalian ruang")
    lines.append("    kunci ketiga algoritma). IC dan entropinya setara dengan")
    lines.append("    Vigenere + Hill karena transposisi tidak mengubah distribusi")
    lines.append("    frekuensi, tetapi lapisan transposisi tetap menambah")
    lines.append("    keamanan terhadap analisis posisi dengan biaya waktu proses")
    lines.append("    yang lebih tinggi dibandingkan algoritma tunggal.")
    lines.append("")
    lines.append("Analisis dispersi posisi (position dispersion):")
    trans_names = [s.name for s in algorithms.SCENARIOS
                   if "Transposisi" in s.components]
    non_trans_names = [s.name for s in algorithms.SCENARIOS
                       if "Transposisi" not in s.components]
    trans_stdev_avg = sum(dispersion[n]["stdev_mean"] for n in trans_names) \
        / len(trans_names)
    non_trans_stdev_avg = sum(dispersion[n]["stdev_mean"]
                              for n in non_trans_names) / len(non_trans_names)
    trans_gap_avg = sum(dispersion[n]["mean_gap_mean"] for n in trans_names) \
        / len(trans_names)
    non_trans_gap_avg = sum(dispersion[n]["mean_gap_mean"]
                            for n in non_trans_names) / len(non_trans_names)

    def _name_for_components(comps):
        for s in algorithms.SCENARIOS:
            if s.components == comps:
                return s.name
        return None

    lines.append("  * stdev posisi: rata-rata DENGAN transposisi = {:.2f} vs "
                 "TANPA = {:.2f}.".format(trans_stdev_avg, non_trans_stdev_avg))
    lines.append("  * mean gap: rata-rata DENGAN transposisi = {:.2f} vs "
                 "TANPA = {:.2f}.".format(trans_gap_avg, non_trans_gap_avg))
    lines.append("  * Perbandingan berpasangan (tanpa vs dengan transposisi):")
    for s in algorithms.SCENARIOS:
        if "Transposisi" not in s.components:
            continue
        base_comps = [c for c in s.components if c != "Transposisi"]
        base_name = _name_for_components(base_comps)
        if base_name is None:
            continue
        bs = dispersion[base_name]["stdev_mean"]
        ts = dispersion[s.name]["stdev_mean"]
        bg = dispersion[base_name]["mean_gap_mean"]
        tg = dispersion[s.name]["mean_gap_mean"]
        verdict = ("LEBIH TINGGI" if ts > bs else
                   ("lebih rendah" if ts < bs else "sama"))
        lines.append("    - {} vs {} : stdev {:.2f} -> {:.2f} ({}); "
                     "mean gap {:.2f} -> {:.2f} ({})".format(
                         base_name, s.name, bs, ts, verdict, bg, tg, verdict))
    lines.append("  * Skenario dengan transposisi: {}.".format(
        ", ".join(trans_names)))
    lines.append("  * Skenario tanpa transposisi: {}.".format(
        ", ".join(non_trans_names)))
    lines.append("  * Catatan: dispersi hanya naik pada pasangan yang menyertakan")
    lines.append("    Hill. Transposisi tunggal dan Vigenere+Transposisi tetap 0")
    lines.append("    karena membalik 1 karakter plaintext hanya mengubah 1 posisi")
    lines.append("    ciphertext (tidak ada >1 posisi untuk 'disebarkan'). Dengan")
    lines.append("    kata lain transposisi memperbesar sebaran HANYA bila sudah")
    lines.append("    ada banyak posisi yang berubah (yang disediakan oleh Hill).")
    lines.append("")
    lines.append("Catatan ruang kunci gabungan:")
    lines.append("  * Perkalian ruang kunci tiap lapisan adalah ukuran teoretis")
    lines.append("    terhadap brute force naif dan TIDAK menjamin ketahanan")
    lines.append("    terhadap serangan canggih seperti meet-in-the-middle attack.")
    lines.append("")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Titik masuk
# ---------------------------------------------------------------------------
def main():
    print()
    print("#" * 70)
    print("#  ANALISIS KINERJA & KEAMANAN MULTIPLE SUPER ENCRYPTION       #")
    print("#  Vigenere + Hill + Transposisi Kolom (7 skenario)            #")
    print("#" * 70)
    print()

    # 1. Correctness
    run_correctness()

    # Siapkan plaintext uji dari dokumen resmi negara (sama untuk semua
    # skenario), dinormalkan menjadi huruf A-Z.  Untuk pengukuran waktu dipakai
    # beberapa sampel berbeda per ukuran; untuk metrik statistik dipakai sampel
    # pertama tiap ukuran.
    plaintext_sets = sample_data.build_sample_sets(SIZES, SAMPLES_PER_SIZE)
    plaintexts = {size: samples[0] for size, samples in plaintext_sets.items()}
    plaintext_metrics = {
        "ic": ev.index_of_coincidence(plaintexts[REFERENCE_SIZE]),
        "entropy": ev.shannon_entropy(plaintexts[REFERENCE_SIZE]),
    }
    print("Plaintext acuan ({} karakter) dari dokumen resmi negara:".format(
        REFERENCE_SIZE))
    print("  IC {:.4f} | entropi {:.4f} bit (IC acak ideal {:.4f}, entropi "
          "maksimum {:.4f} bit)".format(
              plaintext_metrics["ic"], plaintext_metrics["entropy"],
              ev.IC_RANDOM, ev.MAX_ENTROPY))
    print()

    # 2. Waktu proses
    time_results = run_time_benchmarks(plaintext_sets)

    # 3. Metrik statistik
    stats = run_statistical_metrics(plaintexts)

    # 3b. Avalanche effect
    avalanche = run_avalanche(plaintexts)

    # 3c. Dispersi posisi
    dispersion = run_dispersion(plaintexts)

    # 4. Ruang kunci
    key_spaces = run_key_space()

    # Susun tabel ringkasan
    rows = _scenario_rows(time_results, stats, key_spaces, plaintext_metrics,
                          avalanche, dispersion)

    # Ekspor
    print("=" * 70)
    print("[5] Ekspor hasil (CSV/Excel) & grafik")
    print("=" * 70)
    export_csvs(rows, time_results, stats, key_spaces, avalanche, dispersion)
    export_excel(rows, time_results, stats, key_spaces, avalanche, dispersion)
    make_plots(rows, time_results, stats, key_spaces, plaintext_metrics,
               avalanche, dispersion)
    print()

    # Ringkasan naratif
    summary = build_summary(rows, time_results, stats, key_spaces,
                            plaintext_metrics, avalanche, dispersion)
    print(summary)
    with open(os.path.join(OUTPUT_DIR, "ringkasan_naratif.txt"), "w") as f:
        f.write(summary + "\n")

    # Cetak tabel ringkasan sederhana
    print("=" * 70)
    print("TABEL RINGKASAN (ukuran {} karakter)".format(REFERENCE_SIZE))
    print("=" * 70)
    header = ("{:<30}{:>10}{:>12}{:>10}{:>10}{:>10}{:>22}".format(
        "Skenario", "Enc(ms)", "Dec(ms)", "IC", "Entropi",
        "Aval. P(%)", "Ruang Kunci"))
    print(header)
    print("-" * len(header))
    for r in rows:
        print("{:<30}{:>10.4f}{:>12.4f}{:>10.4f}{:>10.4f}{:>10.2f}{:>22}".format(
            r["Skenario"], r["Waktu Enkripsi (ms)"], r["Waktu Dekripsi (ms)"],
            r["IC Ciphertext"], r["Entropi Ciphertext (bit)"],
            r["Avalanche Plaintext (%)"], r["Estimasi Brute Force"]))
    print()


if __name__ == "__main__":
    main()
